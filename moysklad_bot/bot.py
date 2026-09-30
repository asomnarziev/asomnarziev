import os
import json
import html
import math
import time
import difflib
import calendar
import datetime as dt
from io import BytesIO
from zoneinfo import ZoneInfo

import requests
import telebot
from telebot import types
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

# --- SOZLAMALAR ---
# Tokenlar kodda emas, .env faylida saqlanadi (.env.example ga qarang)
load_dotenv()
MOYSKLAD_TOKEN = os.environ["MOYSKLAD_TOKEN"]
TELEGRAM_BOT_TOKEN = os.environ["TELEGRAM_BOT_TOKEN"]
MOYSKLAD_API = "https://api.moysklad.ru/api/remap/1.2"
# "Bugun", "Kecha" shu vaqt zonasi bo'yicha hisoblanadi (server UTC da bo'lsa ham)
TIMEZONE = ZoneInfo(os.environ.get("TIMEZONE", "Asia/Tashkent"))

# --- ADMIN VA RUXSATLAR ---
ADMIN_ID = int(os.environ["ADMIN_ID"])
USERS_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "allowed_users.json")


def load_allowed_users():
    try:
        with open(USERS_FILE, encoding="utf-8") as f:
            return set(json.load(f)) | {ADMIN_ID}
    except (FileNotFoundError, ValueError):
        return {ADMIN_ID}


def save_allowed_users():
    with open(USERS_FILE, "w", encoding="utf-8") as f:
        json.dump(sorted(ALLOWED_USERS), f)


# Ruxsat berilganlar faylda saqlanadi, bot qayta ishga tushsa ham yo'qolmaydi
ALLOWED_USERS = load_allowed_users()

bot = telebot.TeleBot(TELEGRAM_BOT_TOKEN, parse_mode="HTML")
headers = {"Authorization": f"Bearer {MOYSKLAD_TOKEN}", "Content-Type": "application/json"}
user_steps = {}


# --- YORDAMCHI FUNKSIYALAR ---
def esc(text):
    """Tovar nomlaridagi <, >, & belgilari HTML xabarni buzmasligi uchun."""
    return html.escape(str(text))


RETRY_STATUSES = {429, 500, 502, 503, 504}
RETRY_DELAYS = [3, 6, 12]  # soniya


def ms_get(path, params):
    """MoySklad band bo'lsa (503 va h.k.) biroz kutib, qayta urinib ko'radi."""
    for delay in RETRY_DELAYS + [None]:
        try:
            r = requests.get(MOYSKLAD_API + path, headers=headers, params=params, timeout=90)
        except (requests.ConnectionError, requests.Timeout):
            if delay is None:
                raise
        else:
            if r.status_code not in RETRY_STATUSES or delay is None:
                r.raise_for_status()
                return r.json()
        time.sleep(delay)


def ms_rows(path, params=None):
    """MoySklad'dan barcha qatorlarni oladi (1000 tadan ko'p bo'lsa sahifalab)."""
    rows, offset = [], 0
    while True:
        data = ms_get(path, dict(params or {}, limit=1000, offset=offset))
        batch = data.get('rows', [])
        rows += batch
        offset += len(batch)
        if not batch or offset >= data.get('meta', {}).get('size', 0):
            return rows


def send_long(chat_id, text):
    """Telegram 4096 belgidan uzun xabarni qabul qilmaydi - qatorlar bo'yicha bo'lib yuboramiz."""
    chunk = ""
    for line in text.split("\n"):
        if len(chunk) + len(line) + 1 > 4000:
            bot.send_message(chat_id, chunk)
            chunk = ""
        chunk += line + "\n"
    if chunk.strip():
        bot.send_message(chat_id, chunk)


# --- ASOSIY MENYU ---
def main_menu(chat_id):
    markup = types.ReplyKeyboardMarkup(resize_keyboard=True)
    markup.add(types.KeyboardButton("📊 Umumiy Hisobot"), types.KeyboardButton("📦 Sotuv Tovarlar Bo'yicha"))
    markup.add(types.KeyboardButton("📁 Tovar Qoldiqlari"), types.KeyboardButton("🔄 Aylanma"))
    bot.send_message(chat_id, "👋 Bo'limni tanlang:", reply_markup=markup)


@bot.message_handler(commands=['start'])
def start(message):
    uid = message.from_user.id
    if uid in ALLOWED_USERS:
        user_steps[message.chat.id] = {}
        main_menu(message.chat.id)
    else:
        username = f"@{message.from_user.username}" if message.from_user.username else "username yo'q"
        info = f"👤 Foydalanuvchi: {esc(message.from_user.first_name)} ({esc(username)})\nID: <code>{uid}</code>"
        markup = types.InlineKeyboardMarkup()
        markup.add(types.InlineKeyboardButton("✅ Ruxsat berish", callback_data=f"auth_allow_{uid}"),
                   types.InlineKeyboardButton("❌ Rad etish", callback_data=f"auth_deny_{uid}"))
        try:
            bot.send_message(ADMIN_ID, f"🔔 <b>Yangi kirish so'rovi:</b>\n{info}", reply_markup=markup)
            bot.send_message(message.chat.id, "⏳ Kirish uchun so'rov yuborildi. Admin tasdiqlashini kuting...")
        except telebot.apihelper.ApiTelegramException:
            bot.send_message(message.chat.id, "⚠️ Admin botni ishga tushirishi kerak.")


@bot.callback_query_handler(func=lambda call: True)
def handle_callbacks(call):
    # Tugmadagi "soat" belgisini o'chiradi
    try:
        bot.answer_callback_query(call.id)
    except telebot.apihelper.ApiTelegramException:
        pass
    try:
        dispatch_callback(call)
    except requests.RequestException as e:
        bot.send_message(call.message.chat.id, f"⚠️ MoySklad bilan bog'lanishda xato:\n<code>{esc(e)}</code>")


def dispatch_callback(call):
    chat_id = call.message.chat.id
    uid = call.from_user.id

    if call.data.startswith("auth_"):
        if uid != ADMIN_ID: return
        action, target_id = call.data.split("_")[1], int(call.data.split("_")[2])
        if action == "allow":
            ALLOWED_USERS.add(target_id)
            save_allowed_users()
            bot.send_message(target_id, "✅ Admin ruxsat berdi! /start bosing.")
            bot.edit_message_text(f"✅ {target_id} ga ruxsat berildi.", chat_id, call.message.message_id)
        else:
            bot.send_message(target_id, "❌ So'rovingiz rad etildi.")
            bot.edit_message_text(f"❌ {target_id} rad etildi.", chat_id, call.message.message_id)
        return

    if uid not in ALLOWED_USERS: return
    if chat_id not in user_steps: user_steps[chat_id] = {}

    if call.data.startswith("cal:"): handle_calendar(call)
    elif call.data.startswith("pr:"): handle_preset(call)
    # Eski xabarlardagi tugmalar ham ishlashi uchun
    elif call.data in ("date_today", "date_yesterday"):
        call.data = "pr:" + call.data[5:]
        handle_preset(call)
    elif call.data == "date_calendar": open_calendar(call, 'r')
    elif call.data == "fdate_calendar": open_calendar(call, 's')
    elif call.data == "fdate_now":
        user_steps[chat_id].update({'moment': None, 'label': "Hozirgi"})
        show_stock_folders(chat_id, message_id=call.message.message_id)
    elif call.data == "back_root": show_stock_folders(chat_id, None, call.message.message_id)
    elif call.data.startswith('nav_'): show_stock_folders(chat_id, call.data.split('_')[1], call.message.message_id)
    elif call.data.startswith('fcalc_'): calculate_folder_stock(call)
    elif call.data.startswith('tx:'): send_turnover_excel(call)


# --- HISOBOTLAR ---
def process_reports(chat_id, s, e, l):
    rtype = user_steps[chat_id].get('report_type')
    bot.send_message(chat_id, f"⏳ <b>{l}</b> hisoboti tayyorlanmoqda...")
    if rtype == "📦 Sotuv Tovarlar Bo'yicha":
        generate_product_sales_report(chat_id, s, e, l)
    elif rtype == "🔄 Aylanma":
        generate_turnover_report(chat_id, s, e, l)
    else:
        generate_final_summary(chat_id, s, e, l)


def generate_product_sales_report(chat_id, s, e, l):
    res = ms_rows("/report/profit/byvariant", {"momentFrom": s, "momentTo": e})
    if not res:
        bot.send_message(chat_id, f"📦 <b>{l}</b> davrida sotuvlar topilmadi.")
        return
    # Eng ko'p sotilganlar birinchi
    res.sort(key=lambda r: r.get('sellSum', 0), reverse=True)
    txt = f"📦 <b>SOTUV TOVARLAR ({l})</b>\n\n"
    for r in res[:25]:
        name = r.get('variantName') or r.get('name')
        if not name: name = r.get('assortment', {}).get('name', "Noma'lum tovar")
        txt += f"🔹 {esc(name[:30])}\n    └ {r.get('sellQuantity', 0):,.0f} ta | ${r.get('sellSum', 0)/100:,.2f}\n"
    if len(res) > 25:
        txt += f"\n... va yana {len(res) - 25} ta tovar"
    send_long(chat_id, txt)


# MoySklad'ning "Прибыли и убытки" hisobotidagi kabi bu moddalar operatsion xarajat hisoblanmaydi
NON_OPERATING_EXPENSE_ITEMS = {"закупка товаров", "возврат", "перемещение", "налоги и сборы"}


def fmt_money(amount, iso):
    if iso == "USD":
        return f"${amount:,.2f}"
    if iso == "UZS":
        return f"{amount:,.0f} so'm"
    return f"{amount:,.2f} {iso}"


class Currencies:
    """Hujjatlar turli valyutada bo'lishi mumkin (masalan, chakana savdo so'mda, hisobot esa dollarda).
    Har bir hujjat summasi MoySklad'da shu hujjatga kiritilgan kurs bo'yicha asosiy valyutaga o'giriladi."""

    def __init__(self):
        rows = ms_rows("/entity/currency")
        self.by_id = {c['id']: c for c in rows}
        base = next((c for c in rows if c.get('default')), {})
        self.base = base.get('isoCode') or "USD"

    def factor(self, doc):
        """1 birlik hujjat valyutasi = factor birlik asosiy valyuta."""
        cur_id = doc.get('rate', {}).get('currency', {}).get('meta', {}).get('href', '').split('/')[-1]
        c = self.by_id.get(cur_id)
        if not c or c.get('default'):
            return 1.0
        m = c.get('multiplicity') or 1
        cur_rate = c.get('rate') or 0
        # Valyuta sozlamasidagi joriy kurs (to'g'ri yoki teskari kurs bo'lishi mumkin)
        expected = (m / cur_rate if c.get('indirect') else cur_rate / m) if cur_rate else None
        v = doc.get('rate', {}).get('value') or cur_rate
        if not v:
            return 1.0
        # Hujjatdagi kurs qaysi yo'nalishda saqlanganini joriy kursga eng yaqin variant bo'yicha aniqlaymiz
        candidates = [v, 1 / v, v / m, m / v]
        if expected:
            return min(candidates, key=lambda f: abs(math.log(f / expected)))
        return v / m

    def to_base(self, doc, *fields):
        return sum(doc.get(f, 0) or 0 for f in fields) / 100 * self.factor(doc)

    def total(self, docs, *fields):
        return sum(self.to_base(d, *fields) for d in docs)


def operating_expenses(payments, cur):
    """Operatsion xarajatlar: o'tkazilgan (проведённые) chiqim to'lovlari, tovar xaridi, qaytarish,
    ko'chirish va soliqlardan tashqari."""
    excluded_ids = {i['id'] for i in ms_rows("/entity/expenseitem")
                    if i.get('name', '').strip().lower() in NON_OPERATING_EXPENSE_ITEMS}
    total = 0
    for r in payments:
        if not r.get('applicable', True):
            continue
        item_id = r.get('expenseItem', {}).get('meta', {}).get('href', '').split('/')[-1]
        if item_id not in excluded_ids:
            total += cur.to_base(r, 'sum')
    return total


def payment_type(doc):
    """Qo'lda ochilgan kirim orderidagi "To'lov turi" qo'shimcha maydoni: 'cash' (Naxt), 'card' (Karta) yoki None."""
    for a in doc.get('attributes', []):
        if normalize(a.get('name')) != "tolov turi":
            continue
        value = a.get('value')
        value = normalize(value.get('name') if isinstance(value, dict) else value)
        if value.startswith("kart"):
            return 'card'
        if value.startswith(("nax", "naq", "nal")):
            return 'cash'
    return None


def linked_to_shift(doc):
    """Kassa smenasiga (розничная смена) bog'langan order - chakana savdo bilan avtomatik yaratilgan, u allaqachon
    chakana savdo summasida bor. Bog'liqlik smena maydonida yoki bog'langan hujjatlarda (operations) bo'lishi mumkin."""
    if doc.get('retailShift'):
        return True
    retail_types = {'retailshift', 'retaildemand', 'retailsalesreturn'}
    return any(op.get('meta', {}).get('type') in retail_types for op in doc.get('operations', []))


def generate_final_summary(chat_id, s, e, l):
    cur = Currencies()
    money = lambda amount: fmt_money(amount, cur.base)

    # 1. Savdo va Yalpi foyda (валовая прибыль) - asosiy valyutada
    p_res = ms_rows("/report/profit/byvariant", {"momentFrom": s, "momentTo": e})
    total_s = sum(r.get('sellSum', 0) for r in p_res) / 100
    gross_p = sum(r.get('profit', 0) for r in p_res) / 100

    # 2. Kassa Orderlari (Kirim va Chiqim)
    period = {"filter": f"moment>={s};moment<={e}"}
    cashin_rows = ms_rows("/entity/cashin", period)
    cashout_rows = ms_rows("/entity/cashout", period)

    # 3. Kassa (real) = o'tkazilgan chakana savdolar (розничные продажи) + qo'lda ochilgan kirim orderlari, to'lov usullari bo'yicha
    retail = [r for r in ms_rows("/entity/retaildemand", period) if r.get('applicable', True)]
    # Qarz = "Сумма из аванса": chekning naqd, karta, QR va oldindan to'lovdan (предоплата) tashqari qismi
    for r in retail:
        paid = sum(r.get(f, 0) or 0 for f in ('cashSum', 'noCashSum', 'qrSum',
                                              'prepaymentCashSum', 'prepaymentNoCashSum', 'prepaymentQrSum'))
        r['_debt'] = max(0, r.get('sum', 0) - paid)
    # Qo'lda ochilgan kirim orderlari (smenaga bog'lanmagan приходный ордер) "To'lov turi" bo'yicha naqd yoki kartaga qo'shiladi
    applied_cashin = [r for r in cashin_rows if r.get('applicable', True) and not linked_to_shift(r)]
    cashin_cash = cur.total([r for r in applied_cashin if payment_type(r) == 'cash'], 'sum')
    cashin_card = cur.total([r for r in applied_cashin if payment_type(r) == 'card'], 'sum')
    kassa_cash = cur.total(retail, 'cashSum') + cashin_cash
    kassa_card = cur.total(retail, 'noCashSum') + cashin_card
    # Qarz kassaga tushgan pul emas: u KASSA tagida ko'rsatiladi, lekin jamiga qo'shilmaydi (KASSA = naqd + karta)
    kassa_total = kassa_cash + kassa_card

    # 4. Foyda = operatsion foyda (операционная прибыль): yalpi foyda - operatsion xarajatlar (kassa + bank to'lovlari)
    operating_p = gross_p - operating_expenses(cashout_rows + ms_rows("/entity/paymentout", period), cur)

    report = (f"🗓 <b>UMUMIY HISOBOT: {l}</b>\n"
              f"━━━━━━━━━━━━━━━━━━━━\n"
              f"💰 SAVDO (Umumiy): {money(total_s)}\n"
              f"📥 KIRIM PULLAR: {money(cur.total(cashin_rows, 'sum'))}\n"
              f"📉 RASXODLAR: {money(cur.total(cashout_rows, 'sum'))}\n"
              f"🏦 <b>KASSA (REAL): {money(kassa_total)}</b>\n"
              f"    ├ 💵 Naqd: {money(kassa_cash)}\n"
              f"    ├ 💳 Karta: {money(kassa_card)}\n"
              f"    └ 📝 Qarz: {money(cur.total(retail, '_debt'))}\n"
              f"━━━━━━━━━━━━━━━━━━━━\n"
              f"💸 <b>FOYDA: {money(operating_p)}</b>")

    bot.send_message(chat_id, report)


# --- AYLANMA (ОБОРОТЫ) ---
TURNOVER_PARTS = [("onPeriodStart", "Boshida"), ("income", "Kirim"), ("outcome", "Chiqim"), ("onPeriodEnd", "Oxirida")]


def fmt_qty(q):
    return f"{q:,.0f}" if float(q).is_integer() else f"{q:,.2f}"


def turnover_rows(s, e):
    """MoySklad "Обороты" hisoboti: har bir tovar bo'yicha davr boshidagi, kirim, chiqim va oxiridagi soni/summasi."""
    return ms_rows("/report/turnover/all", {"momentFrom": s, "momentTo": e})


def turnover_totals(rows):
    return {key: (sum(r.get(key, {}).get('quantity', 0) or 0 for r in rows),
                  sum(r.get(key, {}).get('sum', 0) or 0 for r in rows) / 100) for key, _ in TURNOVER_PARTS}


def generate_turnover_report(chat_id, s, e, l):
    rows = turnover_rows(s, e)
    if not rows:
        bot.send_message(chat_id, f"🔄 <b>{l}</b> davrida tovar harakati topilmadi.")
        return
    base = Currencies().base  # hisobot summalari asosiy valyutada
    money = lambda amount: fmt_money(amount, base)
    totals = turnover_totals(rows)
    icons = {"onPeriodStart": "📦", "income": "📥", "outcome": "📤", "onPeriodEnd": "📦"}
    txt = f"🔄 <b>AYLANMA: {l}</b>\n━━━━━━━━━━━━━━━━━━━━\n"
    for key, name in TURNOVER_PARTS:
        qty, total = totals[key]
        txt += f"{icons[key]} {name}: <b>{fmt_qty(qty)}</b> ta | {money(total)}\n"
    txt += f"━━━━━━━━━━━━━━━━━━━━\n🔢 Tovarlar soni: {len(rows)}\n\n<b>Eng ko'p chiqim bo'lganlar:</b>\n"
    top = sorted(rows, key=lambda r: r.get('outcome', {}).get('quantity', 0) or 0, reverse=True)[:20]
    for r in top:
        a = r.get('assortment', {})
        q = {key: r.get(key, {}).get('quantity', 0) or 0 for key, _ in TURNOVER_PARTS}
        end_sum = (r.get('onPeriodEnd', {}).get('sum', 0) or 0) / 100
        name = a.get('name') or "Noma'lum"
        txt += (f"🔹 {esc(name[:35])}\n"
                f"    {fmt_qty(q['onPeriodStart'])} ➜ +{fmt_qty(q['income'])} / −{fmt_qty(q['outcome'])} ➜ "
                f"<b>{fmt_qty(q['onPeriodEnd'])}</b> ta | {money(end_sum)}\n")
    markup = types.InlineKeyboardMarkup()
    markup.add(types.InlineKeyboardButton("📄 Excel'da to'liq", callback_data=f"tx:{s[:10]}:{e[:10]}"))
    send_long(chat_id, txt)
    bot.send_message(chat_id, f"Barcha {len(rows)} ta tovar jadvali:", reply_markup=markup)


def send_turnover_excel(call):
    chat_id = call.message.chat.id
    _, d1, d2 = call.data.split(":")
    start, end = dt.date.fromisoformat(d1), dt.date.fromisoformat(d2)
    rows = sorted(turnover_rows(f"{d1} 00:00:00", f"{d2} 23:59:59"), key=lambda r: r.get('assortment', {}).get('name', ''))

    wb = Workbook()
    ws = wb.active
    ws.title = "Aylanma"
    bold, center = Font(bold=True), Alignment(horizontal="center", vertical="center")
    fill = PatternFill("solid", fgColor="DDEBF7")
    ws.append([f"Aylanma: {fmt_period(start, end)}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Nomi", "Kod", "Artikul", "O'lch. birl."] + [n for _, n in TURNOVER_PARTS for _ in (0, 1)])
    ws.append(["", "", "", ""] + ["Soni", "Summa"] * len(TURNOVER_PARTS))
    for col in range(1, 5):
        ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)
    for i in range(len(TURNOVER_PARTS)):
        ws.merge_cells(start_row=2, start_column=5 + 2 * i, end_row=2, end_column=6 + 2 * i)
    for row in ws.iter_rows(min_row=2, max_row=3):
        for c in row:
            c.font, c.alignment, c.fill = bold, center, fill
    for r in rows:
        a = r.get('assortment', {})
        values = []
        for key, _ in TURNOVER_PARTS:
            part = r.get(key, {})
            values += [part.get('quantity', 0) or 0, (part.get('sum', 0) or 0) / 100]
        ws.append([a.get('name', ''), a.get('code', ''), a.get('article', ''), a.get('uom', {}).get('name', '')] + values)
    totals = turnover_totals(rows)
    ws.append(["JAMI", "", "", ""] + [v for key, _ in TURNOVER_PARTS for v in totals[key]])
    for c in ws[ws.max_row]:
        c.font = bold
    for row in ws.iter_rows(min_row=4, min_col=5):
        for c in row:
            is_sum = (c.column - 5) % 2
            c.number_format = '#,##0.00' if is_sum or not float(c.value or 0).is_integer() else '#,##0'
    ws.column_dimensions["A"].width = 45
    for col, width in zip("BCDEFGHIJKL", [10, 12, 10] + [10, 14] * 4):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "B4"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = f"aylanma_{start:%d.%m.%Y}-{end:%d.%m.%Y}.xlsx"
    bot.send_document(chat_id, buf, caption=f"🔄 Aylanma: {fmt_period(start, end)} ({len(rows)} ta tovar)")


# --- QOLDIQLAR ---
def parent_folder_id(folder):
    return folder.get('productFolder', {}).get('meta', {}).get('href', '').split('/')[-1]


def show_stock_folders(chat_id, parent_id=None, message_id=None):
    all_f = ms_rows("/entity/productfolder")
    curr = [f for f in all_f if parent_folder_id(f) == parent_id] if parent_id else [f for f in all_f if 'productFolder' not in f]
    markup = types.InlineKeyboardMarkup(row_width=1)
    for f in curr:
        f_id = f['id']
        has_sub = any(parent_folder_id(sf) == f_id for sf in all_f)
        markup.add(types.InlineKeyboardButton(text=f"📁 {f['name']}", callback_data=f"{'nav' if has_sub else 'fcalc'}_{f_id}"))
    if parent_id:
        markup.add(types.InlineKeyboardButton(text="✅ Hisoblash", callback_data=f"fcalc_{parent_id}"),
                   types.InlineKeyboardButton(text="⬅️ Orqaga", callback_data="back_root"))
    title = "📂 Bo'limni tanlang:"
    if message_id: bot.edit_message_text(title, chat_id, message_id, reply_markup=markup)
    else: bot.send_message(chat_id, title, reply_markup=markup)


def calculate_folder_stock(call):
    chat_id, f_id = call.message.chat.id, call.data.split('_')[1]
    moment = user_steps[chat_id].get('moment')
    params = {"filter": f"productFolder={MOYSKLAD_API}/entity/productfolder/{f_id}", "includeEmpty": "true"}
    if moment: params["moment"] = moment
    data = ms_rows("/report/stock/all", params)
    qty, total, details = 0, 0, ""
    for r in data:
        stock = r.get('stock', 0)
        if stock > 0:
            price = (r.get('price', 0) or 0) / 100
            qty += stock
            total += (stock * price)
            name = r.get('name', "Noma'lum")[:25]
            details += f"🔹 {esc(name)}: {stock:,.0f} ta | ${stock*price:,.2f}\n"
    label = user_steps[chat_id].get('label', 'Hozirgi')
    send_long(chat_id, f"📋 <b>QOLDIQLAR ({label})</b>\n\n{details}\n🔢 Jami: {qty:,.0f} ta\n💰 Qiymati: ${total:,.2f}")


@bot.message_handler(func=lambda m: m.from_user.id in ALLOWED_USERS and m.text in ["📊 Umumiy Hisobot", "📦 Sotuv Tovarlar Bo'yicha", "🔄 Aylanma"])
def sales_init(message):
    user_steps[message.chat.id] = {'report_type': message.text}
    btn = types.InlineKeyboardButton
    markup = types.InlineKeyboardMarkup()
    markup.row(btn("Bugun", callback_data="pr:today"), btn("Kecha", callback_data="pr:yesterday"))
    markup.row(btn("Shu hafta", callback_data="pr:week"), btn("O'tgan hafta", callback_data="pr:lastweek"))
    markup.row(btn("Shu oy", callback_data="pr:month"), btn("O'tgan oy", callback_data="pr:lastmonth"))
    markup.row(btn("📅 Kalendardan tanlash", callback_data="cal:o:r"))
    bot.send_message(message.chat.id, f"<b>{message.text}</b>\n\nDavrni tanlang:", reply_markup=markup)


@bot.message_handler(func=lambda m: m.from_user.id in ALLOWED_USERS and m.text == "📁 Tovar Qoldiqlari")
def stock_report_init(message):
    user_steps[message.chat.id] = {'report_type': 'stock'}
    markup = types.InlineKeyboardMarkup()
    markup.add(types.InlineKeyboardButton("Hozirgi holat", callback_data="fdate_now"),
               types.InlineKeyboardButton("📅 Sana tanlash", callback_data="cal:o:s"))
    bot.send_message(message.chat.id, "Vaqtni tanlang:", reply_markup=markup)


# --- KALENDAR ---
UZ_MONTHS = ["Yanvar", "Fevral", "Mart", "Aprel", "May", "Iyun",
             "Iyul", "Avgust", "Sentabr", "Oktabr", "Noyabr", "Dekabr"]
UZ_WEEKDAYS = ["Du", "Se", "Ch", "Pa", "Ju", "Sh", "Ya"]


def today():
    return dt.datetime.now(TIMEZONE).date()


def fmt_period(start, end):
    return f"{start:%d.%m.%Y}" if start == end else f"{start:%d.%m.%Y} — {end:%d.%m.%Y}"


def run_period(chat_id, start, end, name=None):
    label = f"{name} ({fmt_period(start, end)})" if name else fmt_period(start, end)
    process_reports(chat_id, f"{start:%Y-%m-%d} 00:00:00", f"{end:%Y-%m-%d} 23:59:59", label)


def preset_range(key, t):
    """Tayyor davrlar: (boshlanish, tugash, nomi)."""
    month_start = t.replace(day=1)
    week_start = t - dt.timedelta(days=t.weekday())
    if key == "today":
        return t, t, "BUGUN"
    if key == "yesterday":
        y = t - dt.timedelta(days=1)
        return y, y, "KECHA"
    if key == "week":
        return week_start, t, "SHU HAFTA"
    if key == "lastweek":
        return week_start - dt.timedelta(days=7), week_start - dt.timedelta(days=1), "O'TGAN HAFTA"
    if key == "month":
        return month_start, t, "SHU OY"
    if key == "lastmonth":
        last_end = month_start - dt.timedelta(days=1)
        return last_end.replace(day=1), last_end, "O'TGAN OY"
    return None


def handle_preset(call):
    period = preset_range(call.data[3:], today())
    if period:
        run_period(call.message.chat.id, *period)


def calendar_text(mode, start=None):
    if mode == 's':
        return "📅 <b>Qoldiq sanasini tanlang</b>"
    if not start:
        return "📅 <b>Davrni tanlang</b>\n\n1️⃣ Boshlanish kunini bosing"
    return (f"📅 <b>Davrni tanlang</b>\n\n✅ Boshlanish: <b>{start:%d.%m.%Y}</b>\n"
            f"2️⃣ Tugash kunini bosing\n<i>Bir kunlik hisobot uchun shu kunni yana bir marta bosing</i>")


def build_calendar(year, month, mode, start=None):
    """Oy kunlari jadvali. mode: 'r' - davr (2 ta sana), 's' - bitta sana."""
    btn = types.InlineKeyboardButton
    t = today()
    noop = "cal:x"
    prev_y, prev_m = (year - 1, 12) if month == 1 else (year, month - 1)
    next_y, next_m = (year + 1, 1) if month == 12 else (year, month + 1)
    has_next = (next_y, next_m) <= (t.year, t.month)
    markup = types.InlineKeyboardMarkup()
    markup.row(btn("◀️", callback_data=f"cal:n:{mode}:{prev_y}-{prev_m:02d}"),
               btn(f"{UZ_MONTHS[month - 1]} {year}", callback_data=noop),
               btn("▶️" if has_next else " ", callback_data=f"cal:n:{mode}:{next_y}-{next_m:02d}" if has_next else noop))
    markup.row(*[btn(w, callback_data=noop) for w in UZ_WEEKDAYS])
    for week in calendar.Calendar(firstweekday=0).monthdatescalendar(year, month):
        row = []
        for d in week:
            if d.month != month:
                row.append(btn(" ", callback_data=noop))
            elif d > t:
                row.append(btn("·", callback_data=noop))  # kelajak kunlari tanlanmaydi
            else:
                text = f"✅{d.day}" if d == start else (f"•{d.day}•" if d == t else str(d.day))
                row.append(btn(text, callback_data=f"cal:d:{mode}:{d.isoformat()}"))
        markup.row(*row)
    bottom = [btn("❌ Bekor qilish", callback_data="cal:c")]
    if (year, month) != (t.year, t.month):
        bottom.insert(0, btn("↩️ Joriy oy", callback_data=f"cal:n:{mode}:{t.year}-{t.month:02d}"))
    markup.row(*bottom)
    return markup


def open_calendar(call, mode):
    chat_id = call.message.chat.id
    user_steps[chat_id].pop('range_start', None)
    t = today()
    bot.edit_message_text(calendar_text(mode), chat_id, call.message.message_id,
                          reply_markup=build_calendar(t.year, t.month, mode))


def handle_calendar(call):
    chat_id, msg_id = call.message.chat.id, call.message.message_id
    parts = call.data.split(":")
    action = parts[1]
    if action == "x":
        return
    if action == "c":
        user_steps[chat_id].pop('range_start', None)
        bot.edit_message_text("❌ Bekor qilindi.", chat_id, msg_id)
        return
    if action == "o":
        open_calendar(call, parts[2])
        return
    mode = parts[2]
    start = user_steps[chat_id].get('range_start') if mode == 'r' else None
    if action == "n":
        year, month = map(int, parts[3].split("-"))
        bot.edit_message_text(calendar_text(mode, start), chat_id, msg_id, reply_markup=build_calendar(year, month, mode, start))
        return
    if action != "d":
        return
    picked = dt.date.fromisoformat(parts[3])
    if mode == 's':
        user_steps[chat_id].update({'moment': f"{picked:%Y-%m-%d} 23:59:59", 'label': f"{picked:%d.%m.%Y}"})
        show_stock_folders(chat_id, message_id=msg_id)
    elif not start:
        user_steps[chat_id]['range_start'] = picked
        bot.edit_message_text(calendar_text(mode, picked), chat_id, msg_id,
                              reply_markup=build_calendar(picked.year, picked.month, mode, picked))
    else:
        user_steps[chat_id].pop('range_start', None)
        start, end = sorted((start, picked))
        bot.edit_message_text(f"📅 Davr: <b>{fmt_period(start, end)}</b>", chat_id, msg_id)
        run_period(chat_id, start, end)


# --- TOVAR QIDIRISH ---
MENU_BUTTONS = {"📊 Umumiy Hisobot", "📦 Sotuv Tovarlar Bo'yicha", "📁 Tovar Qoldiqlari", "🔄 Aylanma"}
STOCK_CACHE_SECONDS = 60
_stock_cache = {"time": 0, "rows": []}


def get_all_stock():
    """Barcha tovarlar qoldig'i. Har bir xabarda MoySklad'ga qayta so'rov yubormaslik uchun 1 daqiqa saqlanadi."""
    if time.time() - _stock_cache["time"] > STOCK_CACHE_SECONDS:
        _stock_cache["rows"] = ms_rows("/report/stock/all", {"includeEmpty": "true"})
        _stock_cache["time"] = time.time()
    return _stock_cache["rows"]


# Kirill (rus va o'zbek) harflarini lotinga o'giramiz, shunda "Холодильник" ham "xolodilnik" deb topiladi
CYR_TO_LAT = {
    'а': 'a', 'б': 'b', 'в': 'v', 'г': 'g', 'д': 'd', 'е': 'e', 'ё': 'yo', 'ж': 'j', 'з': 'z',
    'и': 'i', 'й': 'y', 'к': 'k', 'л': 'l', 'м': 'm', 'н': 'n', 'о': 'o', 'п': 'p', 'р': 'r',
    'с': 's', 'т': 't', 'у': 'u', 'ф': 'f', 'х': 'x', 'ц': 'ts', 'ч': 'ch', 'ш': 'sh', 'щ': 'sh',
    'ъ': '', 'ы': 'i', 'ь': '', 'э': 'e', 'ю': 'yu', 'я': 'ya',
    'ў': 'o', 'қ': 'q', 'ғ': 'g', 'ҳ': 'h',
}


def normalize(text):
    """Qidiruv uchun matnni bir xil ko'rinishga keltiradi (faqat solishtirish uchun, foydalanuvchiga ko'rsatilmaydi):
    kichik harf, kirill -> lotin, apostroflarsiz (qo'l = qol), h = x (holodilnik = xolodilnik)."""
    text = "".join(CYR_TO_LAT.get(ch, ch) for ch in str(text or "").lower())
    for ch in "'‘’ʻʼ`´":
        text = text.replace(ch, "")
    text = text.replace("h", "x")
    return " ".join(text.split())


def find_products(query, rows):
    """Avval nomida (yoki artikul/kodida) so'rovdagi barcha so'zlar borlarini qidiradi.
    Hech narsa topilmasa, xato yozilgan so'zlarga yaqin nomlarni qidiradi."""
    q = normalize(query)
    words = q.split()
    exact = []
    for r in rows:
        name = normalize(r.get('name'))
        text = f"{name} {normalize(r.get('article'))} {normalize(r.get('code'))}"
        if all(w in text for w in words):
            exact.append((0 if name.startswith(q) else 1, name, r))
    if exact:
        return [r for *_, r in sorted(exact, key=lambda x: x[:2])], True

    fuzzy = []
    for r in rows:
        name = normalize(r.get('name'))
        name_words = name.split()
        if not name_words:
            continue
        # Har bir so'z uchun nomdagi eng o'xshash so'zni topamiz
        score = sum(max(difflib.SequenceMatcher(None, w, nw).ratio() for nw in name_words) for w in words) / len(words)
        score = max(score, difflib.SequenceMatcher(None, q, name).ratio())
        if score >= 0.75:
            fuzzy.append((-score, name, r))
    return [r for *_, r in sorted(fuzzy, key=lambda x: x[:2])], False


@bot.message_handler(func=lambda m: m.from_user.id in ALLOWED_USERS and m.text and not m.text.startswith("/") and m.text not in MENU_BUTTONS)
def product_search(message):
    query = message.text.strip()
    if len(query) < 2:
        bot.send_message(message.chat.id, "🔎 Kamida 2 ta harf yozing.")
        return
    try:
        rows = get_all_stock()
    except requests.RequestException as e:
        bot.send_message(message.chat.id, f"⚠️ MoySklad bilan bog'lanishda xato:\n<code>{esc(e)}</code>")
        return
    found, is_exact = find_products(query, rows)
    if not found:
        bot.send_message(message.chat.id, f"❌ <b>{esc(query)}</b> — topilmadi.")
        return
    title = "🔎 Natijalar" if is_exact else "🔎 Aniq topilmadi, o'xshash tovarlar"
    txt = f"{title}: <b>{esc(query)}</b>\n\n"
    for r in found[:20]:
        name = r.get('name') or "Noma'lum"
        txt += f"🔹 {esc(name)}\n    └ Qoldiq: <b>{r.get('stock', 0):,.0f}</b> ta\n"
    if len(found) > 20:
        txt += f"\n... va yana {len(found) - 20} ta. Aniqroq yozing."
    send_long(message.chat.id, txt)


if __name__ == "__main__":
    print("Bot ishga tushdi...")
    bot.infinity_polling()
